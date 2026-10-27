"""Stage 5: the dashboard and the finance reports."""
import io

from openpyxl import load_workbook

from app.reports import safe
from app.seed import PLANTED, business_days
from tests.conftest import END

DAYS = business_days(END)


def test_dashboard_for_a_day_matches_the_other_pages(client):
    pos_, code, sar = PLANTED["cash_differences"][0]  # D-103 50 short, and D-105's terminal 150 short on the same day
    d = DAYS[pos_].isoformat()
    dash = client.get("/api/dashboard", params={"date": d}).json()
    f = dash["figures"]
    coll = client.get("/api/collections", params={"date": d}).json()["totals"]
    possum = client.get("/api/pos", params={"date": d}).json()["totals"]
    dep = next(r for r in client.get("/api/deposits", params={"end": END.isoformat(), "count": 30}).json()["rows"] if r["date"] == d)
    assert f["cash_received"] == coll["cash_received"] and f["cash_short"] == 50.0 and f["short_count"] == 1
    assert f["pos_difference"] == possum["difference"] == -150.0 and f["pos_system"] == possum["system_total"]
    assert f["to_deposit"] == dep["to_deposit"] and f["pending"] == dep["pending"]
    assert f["total_collected"] == round(f["cash_received"] + f["pos_terminal"], 2)
    assert dash["previous_date"] == DAYS[pos_ + 1].isoformat() and dash["previous"]["cash_received"] > 0
    issues = {a["driver_code"]: [i["kind"] for i in a["issues"]] for a in dash["attention"]}
    assert issues["D-103"] == ["cash_short"] and issues["D-105"] == ["terminal_short"]
    assert len(dash["trend"]) == 30 and dash["trend"][-1]["date"] == d


def test_dashboard_alerts_cover_every_kind_of_open_item(client):
    dash = client.get("/api/dashboard", params={"date": END.isoformat()}).json()
    kinds = dash["alert_kinds"]
    assert kinds["deposit_short"] >= 1 and kinds["pos_open"] >= 1
    assert sum(kinds.values()) == dash["alert_count"] and len(dash["alerts"]) == min(12, dash["alert_count"])
    dates = [a["date"] for a in dash["alerts"]]
    assert dates == sorted(dates, reverse=True)  # newest first
    assert dash["cash_on_hand"] > 0


def test_safe_text_for_excel():
    assert safe("=SUM(A1)") == "'=SUM(A1)" and safe("@cmd") == "'@cmd" and safe("Fuel") == "Fuel" and safe(None) == ""


def test_daily_excel_report(client):
    pos_, code, _ = PLANTED["pos_differences"][0]
    d = DAYS[pos_].isoformat()
    r = client.get("/api/reports/day", params={"date": d, "format": "xlsx"})
    assert r.status_code == 200 and f"codesk-daily-{d}.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "By day", "Collections", "POS check", "Expenses", "Deposits"]
    coll = wb["Collections"]
    assert coll.max_row == 11 + 2  # header, 11 drivers, total
    assert coll.cell(row=13, column=5).value == "=SUM(E2:E12)"  # a real formula, not a typed number
    pos_sheet = [[c.value for c in row] for row in wb["POS check"].iter_rows(min_row=2)]
    short = next(row for row in pos_sheet if row[1] == code)
    assert short[7] == -150.0 and short[8] == "Terminal short"
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=4, max_row=14) if row[0].value}
    day = client.get("/api/dashboard", params={"date": d}).json()["figures"]
    assert summary["Cash received from drivers"] == day["cash_received"] and summary["Cash to deposit"] == day["to_deposit"]


def test_formula_like_text_is_neutralised_in_the_report(client):
    did = next(x["id"] for x in client.get("/api/drivers").json() if x["code"] == "D-110")
    day = "2026-05-05"
    client.put("/api/collections", json={"business_date": day, "driver_id": did, "expected_cash": 10, "cash_received": 10,
                                         "note": "=HYPERLINK(\"http://x\")"})
    wb = load_workbook(io.BytesIO(client.get("/api/reports/day", params={"date": day}).content))
    note = wb["Collections"].cell(row=2, column=10).value
    assert note.startswith("'=")


def test_pdf_reports(client):
    r = client.get("/api/reports/day", params={"date": DAYS[1].isoformat(), "format": "pdf"})
    assert r.status_code == 200 and r.content[:5] == b"%PDF-" and r.headers["content-type"] == "application/pdf"
    r = client.get("/api/reports/range", params={"start": DAYS[29].isoformat(), "end": END.isoformat(), "format": "pdf"})
    assert r.status_code == 200 and r.content[:5] == b"%PDF-" and len(r.content) > 5000


def test_range_report_and_preview(client):
    start, end = DAYS[29].isoformat(), END.isoformat()
    p = client.get("/api/reports/preview", params={"start": start, "end": end}).json()
    assert p["days"] == 30 and p["handovers"] == 330 and p["cash_differences"] == len(PLANTED["cash_differences"])
    assert p["pos_differences"] == len(PLANTED["pos_differences"]) and p["open_advances"] >= 2
    wb = load_workbook(io.BytesIO(client.get("/api/reports/range", params={"start": start, "end": end}).content))
    assert wb["By day"].max_row == 30 + 2
    assert client.get("/api/reports/range", params={"start": end, "end": start}).status_code == 422
    assert client.get("/api/reports/range", params={"start": "2026-01-01", "end": end}).status_code == 422
    assert client.get("/api/reports/day", params={"date": end, "format": "doc"}).status_code == 422
