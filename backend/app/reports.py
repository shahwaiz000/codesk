"""Daily and date-range finance reports, as an Excel workbook or a printable PDF.

Both formats are built from report_data(), which uses the same calculations as the screens, so the
report and the app always agree. Totals in Excel are real SUM formulas, and any text that Excel could
run as a formula (starting with = + - @) is made harmless first."""
import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from xml.sax.saxutils import escape

from . import collections, pos
from .config import settings
from .database import get_db
from .deposits import ledger, by_day
from .deps import current_user, parse_day
from .expenses import CATEGORY_TEXT
from .models import Collection, Deposit, Driver, Expense
from .money import to_sar

router = APIRouter(prefix="/api/reports", tags=["reports"])
MAX_DAYS = 93


def safe(text) -> str:
    s = "" if text is None else str(text)
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s


def report_data(db: Session, first: date, last: date) -> dict:
    deposit_rows = [r for r in ledger(db) if first.isoformat() <= r["date"] <= last.isoformat()]
    dep_by_day = {r["date"]: r for r in by_day(deposit_rows)}
    handovers = (db.query(Collection).join(Driver).filter(Collection.business_date.between(first, last))
                 .order_by(Collection.business_date, Driver.code).all())
    days = sorted({h.business_date for h in handovers} | {date.fromisoformat(r["date"]) for r in deposit_rows})
    pos_rows, summary = [], []
    for d in days:
        rows, totals = pos.day_summary(db, d)
        pos_rows += [{**r, "date": d.isoformat()} for r in rows]
        hs = [h for h in handovers if h.business_date == d]
        dep = dep_by_day.get(d.isoformat(), {})
        summary.append({
            "date": d.isoformat(), "drivers": len(hs),
            "expected_cash": to_sar(sum(h.expected_cash for h in hs)), "cash_received": to_sar(sum(h.cash_received for h in hs)),
            "cash_difference": to_sar(sum(h.cash_received - h.expected_cash for h in hs)),
            "pos_system": totals["system_total"], "pos_terminal": totals["machines_total"], "pos_difference": totals["difference"],
            "paid_out": dep.get("paid_out", 0.0), "returned": dep.get("returned", 0.0), "to_deposit": dep.get("to_deposit", 0.0),
            "deposited": dep.get("deposited", 0.0), "pending": dep.get("pending", 0.0),
            "deposit_status": dep.get("status_text", "Nothing to deposit"), "late": dep.get("late", False)})
    expenses = (db.query(Expense).join(Driver).filter(Expense.business_date.between(first, last))
                .order_by(Expense.business_date, Expense.category == "advance", Driver.code).all())
    slips = db.query(Deposit).filter(Deposit.business_date.between(first, last)).order_by(Deposit.business_date, Deposit.id).all()
    open_advances = (db.query(Expense).join(Driver).filter(Expense.category == "advance", Expense.settled_on.is_(None),
                                                           Expense.business_date <= last).order_by(Expense.business_date).all())
    pending_days = [r for r in ledger(db) if r["pending"] > 0 and r["date"] <= last.isoformat()]

    def total(key):
        return round(sum(s[key] for s in summary), 2)

    return {
        "first": first, "last": last, "summary": summary,
        "totals": {k: total(k) for k in ("expected_cash", "cash_received", "cash_difference", "pos_system", "pos_terminal",
                                         "pos_difference", "paid_out", "returned", "to_deposit", "deposited", "pending")},
        "cash_mode": collections.cash_mode(db),
        "handovers": [{"date": h.business_date.isoformat(), "driver_code": h.driver.code, "driver_name": h.driver.name,
                       "orders": h.orders, "expected_cash": to_sar(h.expected_cash), "cash_received": to_sar(h.cash_received),
                       "difference": to_sar(h.cash_received - h.expected_cash), "pos_terminal": h.pos_terminal,
                       "pos_amount": to_sar(h.pos_amount), "note": h.note} for h in handovers],
        "pos": pos_rows,
        "expenses": [{"date": e.business_date.isoformat(), "driver_code": e.driver.code, "driver_name": e.driver.name,
                      "type": CATEGORY_TEXT[e.category], "amount": to_sar(e.amount), "reference": e.reference, "note": e.note,
                      "recovered": (f"{e.settled_on.isoformat()} ({e.settlement_method})" if e.settled_on else
                                    "Open" if e.category == "advance" else "")} for e in expenses],
        "slips": [{"date": s.business_date.isoformat(), "deposit_date": s.deposit_date.isoformat(), "bank": s.bank,
                   "slip_ref": s.slip_ref, "amount": to_sar(s.amount)} for s in slips],
        "open_advances": [{"date": e.business_date.isoformat(), "driver_code": e.driver.code, "driver_name": e.driver.name,
                           "amount": to_sar(e.amount), "days_open": (last - e.business_date).days} for e in open_advances],
        "pending_days": [{"date": r["date"], "pending": r["pending"], "status": r["status_text"]} for r in pending_days],
    }


def _range(start: str, end: str) -> tuple[date, date]:
    first, last = parse_day(start, "start"), parse_day(end, "end")
    if last < first:
        raise HTTPException(422, "The end date is before the start date")
    if (last - first).days > MAX_DAYS:
        raise HTTPException(422, f"Choose at most {MAX_DAYS} days for one report")
    return first, last


@router.get("/preview")
def preview(start: str, end: str, db: Session = Depends(get_db)):
    first, last = _range(start, end)
    data = report_data(db, first, last)
    return {"start": first.isoformat(), "end": last.isoformat(), "days": len(data["summary"]), "totals": data["totals"],
            "handovers": len(data["handovers"]), "cash_differences": sum(1 for h in data["handovers"] if h["difference"]),
            "pos_differences": sum(1 for p in data["pos"] if p["status"] in ("terminal_short", "terminal_over")),
            "expenses": len(data["expenses"]), "slips": len(data["slips"]),
            "open_advances": len(data["open_advances"]), "pending_days": len(data["pending_days"])}


# ---------------------------------------------------------------- Excel

GREEN, PAPER, LINE = "0F2E22", "F3F5F2", "E1E6E0"
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=GREEN)
TOTAL_FILL = PatternFill("solid", fgColor=PAPER)
THIN = Border(bottom=Side(style="thin", color=LINE))
MONEY = '#,##0.00;[Red]-#,##0.00'


def _sheet(wb: Workbook, title: str, columns: list[tuple[str, str, int]], rows: list[dict], totals: tuple[str, ...] = ()):
    """columns: (key, header, width). Money columns are those listed in totals or ending in _amount style keys."""
    ws = wb.create_sheet(title)
    money_keys = set(totals) | {k for k, _, _ in columns if k in ("amount", "pending", "difference", "pos_amount")}
    for i, (_, header, width) in enumerate(columns, 1):
        c = ws.cell(row=1, column=i, value=header)
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    for r, row in enumerate(rows, 2):
        for i, (key, _, _) in enumerate(columns, 1):
            v = row.get(key)
            c = ws.cell(row=r, column=i, value=v if isinstance(v, (int, float)) and not isinstance(v, bool) else safe(v))
            c.border = THIN
            if key in money_keys or isinstance(v, float):
                c.number_format = MONEY
    if totals and rows:
        tr = len(rows) + 2
        ws.cell(row=tr, column=1, value="Total").font = Font(bold=True)
        for i, (key, _, _) in enumerate(columns, 1):
            ws.cell(row=tr, column=i).fill = TOTAL_FILL
            if key in totals:
                col = get_column_letter(i)
                c = ws.cell(row=tr, column=i, value=f"=SUM({col}2:{col}{tr - 1})")
                c.font, c.number_format = Font(bold=True), MONEY
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    return ws


def build_xlsx(data: dict, title: str, user: str) -> bytes:
    wb = Workbook()
    cover = wb.active
    cover.title = "Summary"
    cover["A1"] = title
    cover["A1"].font = Font(bold=True, size=16, color=GREEN)
    cover["A2"] = f"{data['first'].isoformat()} to {data['last'].isoformat()}  ·  prepared by {user} on {datetime.now():%Y-%m-%d %H:%M}"
    cover["A2"].font = Font(color="7C8781")
    t = data["totals"]
    lines = [("Cash expected (system)", t["expected_cash"]), ("Cash received from drivers", t["cash_received"]),
             ("Cash difference (short is negative)", t["cash_difference"]), ("POS per system orders", t["pos_system"]),
             ("POS per terminals", t["pos_terminal"]), ("POS difference (terminal minus system)", t["pos_difference"]),
             ("Paid out: expenses and advances", t["paid_out"]), ("Advances returned in cash", t["returned"]),
             ("Cash to deposit", t["to_deposit"]), ("Deposited", t["deposited"]), ("Still pending", t["pending"])]
    for i, (label, value) in enumerate(lines, 4):
        cover.cell(row=i, column=1, value=label)
        c = cover.cell(row=i, column=2, value=value)
        c.number_format = MONEY
        if label in ("Cash to deposit", "Still pending"):
            cover.cell(row=i, column=1).font = Font(bold=True)
            c.font = Font(bold=True)
    cover.column_dimensions["A"].width = 42
    cover.column_dimensions["B"].width = 18
    r = len(lines) + 6
    cover.cell(row=r, column=1, value="Open items at the end of the period").font = Font(bold=True)
    cover.cell(row=r + 1, column=1, value="Days with cash still pending")
    cover.cell(row=r + 1, column=2, value=len(data["pending_days"]))
    cover.cell(row=r + 2, column=1, value="Advances not yet recovered")
    c = cover.cell(row=r + 2, column=2, value=round(sum(a["amount"] for a in data["open_advances"]), 2))
    c.number_format = MONEY
    cover.cell(row=r + 4, column=1, value="All figures in SAR. Fictional demo data.").font = Font(italic=True, color="7C8781")

    _sheet(wb, "By day", [("date", "Business day", 13), ("drivers", "Drivers", 9), ("expected_cash", "Cash expected", 15),
                          ("cash_received", "Cash received", 15), ("cash_difference", "Cash difference", 15),
                          ("pos_system", "POS system", 14), ("pos_terminal", "POS terminal", 14), ("pos_difference", "POS difference", 15),
                          ("paid_out", "Paid out", 12), ("returned", "Returned", 11), ("to_deposit", "To deposit", 15),
                          ("deposited", "Deposited", 15), ("pending", "Pending", 13), ("deposit_status", "Deposit status", 18)],
           data["summary"], ("expected_cash", "cash_received", "cash_difference", "pos_system", "pos_terminal", "pos_difference",
                             "paid_out", "returned", "to_deposit", "deposited", "pending"))
    _sheet(wb, "Collections", [("date", "Business day", 13), ("driver_code", "Driver ID", 15), ("driver_name", "Driver", 22),
                               ("orders", "Orders", 8), ("expected_cash", "Cash expected", 15), ("cash_received", "Cash received", 15),
                               ("difference", "Difference", 13), ("pos_terminal", "POS terminal", 13), ("pos_amount", "POS total", 13),
                               ("note", "Note", 34)],
           data["handovers"], ("expected_cash", "cash_received", "difference", "pos_amount"))
    _sheet(wb, "POS check", [("date", "Business day", 13), ("driver_code", "Driver ID", 15), ("driver_name", "Driver", 22),
                             ("terminal", "Terminal", 12), ("order_count", "POS orders", 11), ("system_total", "System total", 14),
                             ("terminal_total", "Terminal total", 14), ("difference", "Difference", 13), ("status_text", "Status", 18),
                             ("review_text", "Review", 40)],
           [{**p, "review_text": (f"{p['review']['status'].replace('_', ' ')} by {p['review']['by']}: {p['review']['note']}"
                                  if p["review"] else "")} for p in data["pos"]],
           ("system_total", "terminal_total", "difference"))
    _sheet(wb, "Expenses", [("date", "Business day", 13), ("driver_code", "Driver ID", 15), ("driver_name", "Driver", 22),
                            ("type", "Type", 13), ("amount", "Amount", 13), ("reference", "Receipt", 13),
                            ("recovered", "Advance recovered", 22), ("note", "Note", 30)], data["expenses"], ("amount",))
    _sheet(wb, "Deposits", [("date", "Business day", 13), ("deposit_date", "Deposited on", 13), ("bank", "Bank", 24),
                            ("slip_ref", "Slip", 14), ("amount", "Amount", 15)], data["slips"], ("amount",))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

def build_pdf(data: dict, title: str, user: str, daily: bool) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    green, ink, soft, line, red = (colors.HexColor(c) for c in ("#0f2e22", "#16211c", "#7c8781", "#e1e6e0", "#b8402e"))
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=green, spaceAfter=3)
    h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=10.5, textColor=ink, spaceBefore=9, spaceAfter=4)
    small = ParagraphStyle("small", fontName="Helvetica", fontSize=8, textColor=soft)
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=8.5, textColor=ink)

    def m(v):
        return "-" if v is None else f"{v:,.2f}"

    def table(head, rows, widths, money_cols=(), total=None, red_cols=()):
        rows = rows or [["Nothing recorded"] + [""] * (len(head) - 1)]
        body_rows = [head] + rows + ([total] if total else [])
        t = Table(body_rows, colWidths=[w * mm for w in widths], repeatRows=1)
        style = [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.5), ("FONT", (0, 1), (-1, -1), "Helvetica", 7.5),
                 ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("BACKGROUND", (0, 0), (-1, 0), green),
                 ("LINEBELOW", (0, 1), (-1, -1), 0.4, line), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                 ("TOPPADDING", (0, 0), (-1, -1), 2.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2)]
        for c in money_cols:
            style.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
        if total:
            style += [("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 7.5), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f3f5f2"))]
        for r, row in enumerate(rows, 1):
            for c in red_cols:
                if isinstance(row[c], str) and row[c].startswith("-") and row[c] != "-":
                    style.append(("TEXTCOLOR", (c, r), (c, r), red))
        t.setStyle(TableStyle(style))
        return t

    t = data["totals"]
    period = data["first"].strftime("%a %d %b %Y") if daily else f"{data['first']:%d %b %Y} to {data['last']:%d %b %Y}"
    story = [Paragraph(title, h1),
             Paragraph(f"{period}  |  prepared by {user} on {datetime.now():%d %b %Y %H:%M}  |  all amounts in SAR", small),
             Spacer(1, 8)]
    figures = [["Cash received from drivers", m(t["cash_received"]), "POS per terminals", m(t["pos_terminal"])],
               ["Cash expected (system)", m(t["expected_cash"]), "POS per system orders", m(t["pos_system"])],
               ["Cash difference", m(t["cash_difference"]), "POS difference", m(t["pos_difference"])],
               ["Paid out (expenses, advances)", m(t["paid_out"]), "Cash to deposit", m(t["to_deposit"])],
               ["Advances returned in cash", m(t["returned"]), "Deposited", m(t["deposited"])],
               ["", "", "Still pending", m(t["pending"])]]
    ft = Table(figures, colWidths=[52 * mm, 34 * mm, 52 * mm, 34 * mm])
    fs = [("FONT", (0, 0), (-1, -1), "Helvetica", 9), ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("ALIGN", (3, 0), (3, -1), "RIGHT"),
          ("LINEBELOW", (0, 0), (-1, -1), 0.4, line), ("FONT", (2, 3), (3, 3), "Helvetica-Bold", 9),
          ("FONT", (2, 5), (3, 5), "Helvetica-Bold", 9), ("BOX", (0, 0), (-1, -1), 0.6, line),
          ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
    for r, key in ((2, "cash_difference"), (2, "pos_difference"), (5, "pending")):
        col = 1 if key == "cash_difference" else 3
        if t[key] < 0 or (key == "pending" and t[key] > 0):
            fs.append(("TEXTCOLOR", (col, r), (col, r), red))
    ft.setStyle(TableStyle(fs))
    story.append(ft)

    if daily:
        story.append(Paragraph("Cash handover by driver", h2))
        code_w = max(16, max((len(h["driver_code"]) for h in data["handovers"]), default=0) * 1.75 + 5)  # long system IDs fit
        name = lambda h: Paragraph(escape(h["driver_name"]), body)  # noqa: E731  long names wrap instead of running over
        orders = str(sum(h["orders"] for h in data["handovers"]))
        if data["cash_mode"] == "paid":  # orders are closed only when paid: nothing was counted, so no difference
            rows = [[h["driver_code"], name(h), str(h["orders"]), m(h["cash_received"]), m(h["pos_amount"]),
                     m(round(h["cash_received"] + h["pos_amount"], 2))] for h in data["handovers"]]
            story.append(table(["Driver", "Name", "Orders", "Cash", "POS terminal", "Total"], rows,
                               [code_w, 176 - code_w - 14 - 3 * 28, 14, 28, 28, 28], (2, 3, 4, 5),
                               ["Total", "", orders, m(t["cash_received"]), m(t["pos_terminal"]),
                                m(round(t["cash_received"] + t["pos_terminal"], 2))]))
        else:
            rows = [[h["driver_code"], name(h), str(h["orders"]), m(h["expected_cash"]), m(h["cash_received"]),
                     m(h["difference"]) if h["difference"] else "-", m(h["pos_amount"])] for h in data["handovers"]]
            story.append(table(["Driver", "Name", "Orders", "Expected", "Counted", "Difference", "POS terminal"], rows,
                               [code_w, 176 - code_w - 14 - 26 * 3 - 22, 14, 26, 26, 22, 26], (2, 3, 4, 5, 6),
                               ["Total", "", orders, m(t["expected_cash"]), m(t["cash_received"]), m(t["cash_difference"]),
                                m(t["pos_terminal"])], red_cols=(5,)))
        diffs = [p for p in data["pos"] if p["status"] in ("terminal_short", "terminal_over", "no_terminal")]
        story.append(Paragraph("POS check: differences", h2))
        if diffs:
            story.append(table(["Driver", "Terminal", "System", "Terminal total", "Difference", "Review"],
                               [[p["driver_code"], p["terminal"] or "-", m(p["system_total"]), m(p["terminal_total"]), m(p["difference"]),
                                 Paragraph(escape(f"{p['review']['status'].replace('_', ' ')}: {p['review']['note']}" if p["review"] else "Not reviewed"), body)]
                                for p in diffs], [code_w, 22, 24, 26, 22, 82 - code_w], (2, 3, 4), red_cols=(4,)))
        else:
            story.append(Paragraph("Every terminal total agrees with the system's POS orders.", body))
    else:
        story.append(Paragraph("By day", h2))
        rows = [[s["date"], m(s["cash_received"]), m(s["cash_difference"]) if s["cash_difference"] else "-",
                 m(s["pos_difference"]) if s["pos_difference"] else "-", m(s["paid_out"]), m(s["to_deposit"]),
                 m(s["deposited"]), m(s["pending"]) if s["pending"] else "-", s["deposit_status"] + (" (late)" if s["late"] else "")]
                for s in data["summary"]]
        story.append(table(["Day", "Cash", "Cash diff", "POS diff", "Paid out", "To deposit", "Deposited", "Pending", "Status"], rows,
                           [18, 20, 16, 16, 16, 21, 21, 18, 30], (1, 2, 3, 4, 5, 6, 7),
                           ["Total", m(t["cash_received"]), m(t["cash_difference"]), m(t["pos_difference"]), m(t["paid_out"]),
                            m(t["to_deposit"]), m(t["deposited"]), m(t["pending"]), ""], red_cols=(2, 3)))

    story.append(Paragraph("Expenses and advances", h2))
    story.append(table(["Day", "Driver", "Type", "Receipt", "Amount", "Recovered"],
                       [[e["date"], Paragraph(escape(f"{e['driver_code']} {e['driver_name']}"), body), e["type"], e["reference"] or "-", m(e["amount"]), e["recovered"] or "-"]
                        for e in data["expenses"]], [20, 70, 22, 18, 20, 26], (4,),
                       ["Total", "", "", "", m(round(sum(e["amount"] for e in data["expenses"]), 2)), ""] if data["expenses"] else None))
    story.append(Paragraph("Bank deposit slips", h2))
    story.append(table(["Business day", "Deposited on", "Bank", "Slip", "Amount"],
                       [[s["date"], s["deposit_date"], s["bank"], s["slip_ref"], m(s["amount"])] for s in data["slips"]],
                       [26, 26, 54, 32, 38], (4,), ["Total", "", "", "", m(t["deposited"])] if data["slips"] else None))
    open_rows = ([["Cash not yet deposited", d["date"], d["status"], m(d["pending"])] for d in data["pending_days"]]
                 + [["Advance not recovered", a["date"], f"{a['driver_code']} {a['driver_name']}, {a['days_open']} days", m(a["amount"])]
                    for a in data["open_advances"]])
    story.append(Paragraph("Open items at the end of the period", h2))
    story.append(table(["Item", "Day", "Detail", "Amount"], open_rows, [44, 24, 70, 38], (3,)))
    sign = Table([["Prepared by", "Checked by", "Approved by"], ["", "", ""], ["Name, signature, date", "Name, signature, date", "Name, signature, date"]],
                 colWidths=[58 * mm] * 3, rowHeights=[12, 22, 10])
    sign.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5), ("FONT", (0, 2), (-1, 2), "Helvetica", 7),
                              ("TEXTCOLOR", (0, 2), (-1, 2), soft), ("LINEBELOW", (0, 1), (-1, 1), 0.6, ink),
                              ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 10)]))
    story += [Spacer(1, 12), KeepTogether([sign])]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(soft)
        canvas.drawString(18 * mm, 10 * mm, f"CODesk  |  {title}  |  {period}" + ("  |  fictional demo data" if settings.seed_demo_data else ""))
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=13 * mm, bottomMargin=16 * mm,
                      title=title, author=user).build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _send(content: bytes, name: str, fmt: str) -> Response:
    media = "application/pdf" if fmt == "pdf" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return Response(content, media_type=media, headers={"Content-Disposition": f'attachment; filename="{name}.{fmt}"'})


def _fmt(value: str) -> str:
    if value not in ("xlsx", "pdf"):
        raise HTTPException(422, "format must be xlsx or pdf")
    return value


@router.get("/day")
def day_report(date: str, format: str = "xlsx", db: Session = Depends(get_db), user: str = Depends(current_user)):  # noqa: A002
    fmt = _fmt(format)
    d = parse_day(date)
    data = report_data(db, d, d)
    title = "Daily finance report"
    body = build_pdf(data, title, user, daily=True) if fmt == "pdf" else build_xlsx(data, title, user)
    return _send(body, f"codesk-daily-{d.isoformat()}", fmt)


@router.get("/range")
def range_report(start: str, end: str, format: str = "xlsx", db: Session = Depends(get_db), user: str = Depends(current_user)):
    fmt = _fmt(format)
    first, last = _range(start, end)
    data = report_data(db, first, last)
    title = "Finance report"
    body = build_pdf(data, title, user, daily=False) if fmt == "pdf" else build_xlsx(data, title, user)
    return _send(body, f"codesk-{first.isoformat()}-to-{last.isoformat()}", fmt)
